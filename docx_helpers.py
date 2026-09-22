"""Minimal docx builder matching the Module 531 house style (Calibri; H1 1A2B38 16pt; H2 0079C0 13pt;
H3 1A2B38 11.5pt; body 11pt; subtitle italic 9pt 5A6B78; arrow lines italic blue 10pt indented;
tables with eaf2f8 header shading). Built on a template docx from the package so styles/numbering exist."""
import copy, os
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TEMPLATE = os.environ.get("DOCX_TEMPLATE")

class Doc:
    def __init__(self, title=None, landscape=False, template=None):
        self.d = Document(template or TEMPLATE)
        body = self.d.element.body
        for el in list(body):
            if el.tag.endswith("}sectPr"): continue
            body.remove(el)
        sec = self.d.sections[0]
        if landscape:
            sec.orientation = WD_ORIENT.LANDSCAPE
            sec.page_width, sec.page_height = Emu(15840*635), Emu(12240*635)
        else:
            sec.orientation = WD_ORIENT.PORTRAIT
            sec.page_width, sec.page_height = Emu(12240*635), Emu(15840*635)
        for m in ("left_margin","right_margin","top_margin","bottom_margin"): setattr(sec, m, Emu(1100*635))
        cp = self.d.core_properties
        cp.author = "Per Scholas Curriculum"; cp.last_modified_by = "Per Scholas Curriculum"
        if title: cp.title = title
        cp.subject = "Module 531: Data Center Logistics Technician Day-One Readiness Immersion"

    # ---- paragraph primitives
    def _p(self, style=None, before=None, after=None, line=None, left=None):
        p = self.d.add_paragraph()
        if style:
            ppr = p._element.get_or_add_pPr(); ps = OxmlElement("w:pStyle"); ps.set(qn("w:val"), style.replace(" ", "")); ppr.insert(0, ps)
        pf = p.paragraph_format
        if before is not None: pf.space_before = Pt(before)
        if after is not None: pf.space_after = Pt(after)
        if line is not None: pf.line_spacing = line
        if left is not None: pf.left_indent = Pt(left)
        return p
    def _r(self, p, text, size=11, bold=False, italic=False, color=None, font="Calibri"):
        r = p.add_run(text); r.font.name = font; r.font.size = Pt(size); r.bold = bold; r.italic = italic
        rpr = r._element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
        if rf is None: rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
        for a in ("w:ascii","w:hAnsi","w:cs","w:eastAsia"): rf.set(qn(a), font)
        if color: r.font.color.rgb = RGBColor.from_string(color)
        return r
    def h1(self, t):  p = self._p("Heading 1", after=4); self._r(p, t, 16, True, color="1A2B38"); return p
    def subtitle(self, t): p = self._p(after=12, line=1.15); self._r(p, t, 9, italic=True, color="5A6B78"); return p
    def h2(self, t):  p = self._p("Heading 2", before=14, after=6); self._r(p, t, 13, True, color="0079C0"); return p
    def h3(self, t):  p = self._p("Heading 3", before=11, after=5); self._r(p, t, 11.5, True, color="1A2B38"); return p
    def body(self, t, bold_lead=None, after=8):
        p = self._p(after=after, line=1.15)
        if bold_lead: self._r(p, bold_lead, 11, True)
        self._runs(p, t); return p
    def _runs(self, p, t, size=11):
        # **bold** inline support
        parts = t.split("**")
        for i, s in enumerate(parts):
            if s: self._r(p, s, size, bold=(i % 2 == 1))
    def bullet(self, t, size=11):
        p = self._p("List Paragraph", after=5, line=1.15)
        ppr = p._element.get_or_add_pPr(); num = OxmlElement("w:numPr")
        il = OxmlElement("w:ilvl"); il.set(qn("w:val"), "0"); ni = OxmlElement("w:numId"); ni.set(qn("w:val"), "2")
        num.append(il); num.append(ni); ppr.insert(0, num)
        self._runs(p, t, size); return p
    def numbered(self, t):
        p = self._p("List Paragraph", after=5, line=1.15)
        ppr = p._element.get_or_add_pPr(); num = OxmlElement("w:numPr")
        il = OxmlElement("w:ilvl"); il.set(qn("w:val"), "0"); ni = OxmlElement("w:numId"); ni.set(qn("w:val"), "1")
        num.append(il); num.append(ni); ppr.insert(0, num)
        self._runs(p, t); return p
    def arrow(self, t):
        p = self._p(after=10, left=18); self._r(p, "→ " + t, 10, italic=True, color="0079C0"); return p
    def banner(self, t, accent="c05000", fill="fbf3ec", border="d9c4b0"):
        p = self._p(after=10); ppr = p._element.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        for side, col, sz in (("top",border,"4"),("left",accent,"24"),("bottom",border,"4"),("right",border,"4")):
            e = OxmlElement(f"w:{side}"); e.set(qn("w:val"),"single"); e.set(qn("w:sz"),sz); e.set(qn("w:space"),"0"); e.set(qn("w:color"),col); bdr.append(e)
        ppr.append(bdr); shd = OxmlElement("w:shd"); shd.set(qn("w:val"),"clear"); shd.set(qn("w:fill"),fill); ppr.append(shd)
        p.paragraph_format.left_indent = Pt(6); p.paragraph_format.right_indent = Pt(6)
        self._runs(p, t); return p
    def note(self, t): return self.banner(t, accent="0079C0", fill="eaf2f8", border="c9d9e6")
    def table(self, header, rows, widths=None, size=9.5, header_fill="eaf2f8"):
        t = self.d.add_table(rows=1 + len(rows), cols=len(header))
        t.style = None
        tblPr = t._element.tblPr
        borders = OxmlElement("w:tblBorders")
        for side in ("top","left","bottom","right","insideH","insideV"):
            e = OxmlElement(f"w:{side}"); e.set(qn("w:val"),"single"); e.set(qn("w:sz"),"4"); e.set(qn("w:space"),"0"); e.set(qn("w:color"),"000000"); borders.append(e)
        tblPr.append(borders)
        lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
        sec = self.d.sections[0]
        avail = (sec.page_width - sec.left_margin - sec.right_margin)
        if widths is None: widths = [1.0/len(header)]*len(header)
        tot = sum(widths); widths = [w/tot for w in widths]
        for ri, row in enumerate([header] + rows):
            for ci, cell in enumerate(row):
                c = t.cell(ri, ci); c.width = Emu(int(avail*widths[ci]))
                tcPr = c._element.get_or_add_tcPr()
                if ri == 0:
                    shd = OxmlElement("w:shd"); shd.set(qn("w:val"),"clear"); shd.set(qn("w:fill"),header_fill); tcPr.append(shd)
                mar = OxmlElement("w:tcMar")
                for side, w in (("top","60"),("left","100"),("bottom","60"),("right","100")):
                    e = OxmlElement(f"w:{side}"); e.set(qn("w:w"), w); e.set(qn("w:type"),"dxa"); mar.append(e)
                tcPr.append(mar)
                p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.05
                txt = str(cell)
                parts = txt.split("**")
                for i, s in enumerate(parts):
                    if s: self._r(p, s, size, bold=(ri == 0 or i % 2 == 1))
        self.d.add_paragraph().paragraph_format.space_after = Pt(4)
        return t
    def page_break(self): self.d.add_page_break()
    def save(self, path): self.d.save(path)
