import os, re, sys, json, glob, subprocess
sys.path.insert(0,'build')
from docx import Document
from pptx import Presentation
import openpyxl
OUT=os.environ["OUT"]
# 1. lockstep: handout questions vs tool LQ
from gen_handouts import parse, HANDOUT_FILES
from spec import LAB_META, LAB_ORDER, tool_file
tot_old=tot_new=0; ok=True
for lid in LAB_ORDER:
    _, items = parse(os.path.join(OUT,"3_Handouts",HANDOUT_FILES[lid]))
    hq=[t for k,t in items if k=="arrow" and ("Answer this as Q" in t or "Attach your image as Q" in t)]
    hkinds=["image" if "Attach" in t else "text" for t in hq]
    hnums=[int(re.search(r"Q(\d+)",t).group(1)) for t in hq]
    steps=[t for k,t in items if k=="body" and re.match(r"Step \d+\.",t)]
    # question prompts in handout = the step paragraph preceding each answer arrow
    prompts=[]
    for i,(k,t) in enumerate(items):
        if k=="arrow" and ("Answer this" in t or "Attach" in t):
            j=i-1
            while items[j][0]!="body": j-=1
            prompts.append(items[j][1])
    for s in LAB_META[lid][3]:
        lq=json.loads(re.search(r'<script id="lqData">var LQ=(.*?);</script>',open(os.path.join(OUT,"4_Lab_Tools",tool_file(lid,s))).read(),re.S).group(1))
        tk=[i["k"] for i in lq["items"]]; tp=[i["p"] for i in lq["items"]]
        good = tk==hkinds and hnums==list(range(1,len(hq)+1)) and tp==prompts
        if not good: ok=False; print("LOCKSTEP FAIL",lid,s,tk,hkinds,hnums, [ (a[:30],b[:30]) for a,b in zip(tp,prompts) if a!=b])
    tot_new+=len(hq)
print("lockstep ok" if ok else "LOCKSTEP PROBLEMS","| total questions",tot_new)
# 2. open all docx/pptx/xlsx
n=0
for f in glob.glob(OUT+"/**/*.docx",recursive=True): Document(f); n+=1
for f in glob.glob(OUT+"/**/*.pptx",recursive=True): Presentation(f); n+=1
for f in glob.glob(OUT+"/**/*.xlsx",recursive=True): openpyxl.load_workbook(f); n+=1
print("opened",n,"office files")
# 3. residual scan
bad=[]
for f in glob.glob(OUT+"/**/*",recursive=True):
    if os.path.isdir(f): continue
    if f.endswith((".html",".md")): txt=open(f,encoding="utf-8").read()
    elif f.endswith(".docx"): txt=subprocess.run(["pandoc",f,"-t","plain","--wrap=none"],capture_output=True,text=True).stdout
    elif f.endswith(".pptx"):
        p=Presentation(f); txt=" ".join(sh.text_frame.text for s in p.slides for sh in s.shapes if sh.has_text_frame)+" ".join(s.notes_slide.notes_text_frame.text for s in p.slides if s.has_notes_slide)
    elif f.endswith(".xlsx"):
        wb=openpyxl.load_workbook(f); txt=" ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value)
    else: continue
    for pat in ["R630","24 kg","Guide(1)","Version 1.2","thirty-kilogram","UCI 2225","2225"]:
        if pat in txt:
            if pat in ("UCI 2225","2225") and "CHANGES.md" in f: continue
            if pat=="Version 1.2" and ("CHANGES" in f): continue
            bad.append((os.path.relpath(f,OUT),pat,txt.count(pat)))
print("residuals:",bad)
# 4. filename cross-reference: every 531-... filename mentioned in docs resolves
names=set(os.path.basename(x) for x in glob.glob(OUT+"/**/*",recursive=True))
missing=set()
for f in glob.glob(OUT+"/**/*.docx",recursive=True)+glob.glob(OUT+"/**/*.md",recursive=True):
    txt=subprocess.run(["pandoc",f,"-t","plain","--wrap=none"],capture_output=True,text=True).stdout if f.endswith(".docx") else open(f).read()
    for m in re.findall(r"531[-_][A-Za-z0-9_\-\.]+?\.(?:html|docx|xlsx|pptx|md)",txt):
        if "[-set]" in m or "[set]" in m: continue
        if m not in names: missing.add((os.path.basename(f),m))
print("missing file refs:",sorted(missing))
