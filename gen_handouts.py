"""Regenerate the handouts: parse the source docx (v1.2, or a v2.0+ handout whose bench additions are stripped
first), add the bench objective, kit equipment, timing, and the bench part whose questions continue the
numbering — same order as the tool's registry."""
import os, re, sys, json
sys.path.insert(0, os.path.dirname(__file__))
from docx import Document
from docx.oxml.ns import qn
import spec
from spec import LABS, LAB_META, LAB_ORDER, tool_files_label, VERSION, VDATE, COURSE, SRC_ID, subst
from docx_helpers import Doc
SRC = os.environ["SRC"]; OUT = os.environ["OUT"]

HANDOUT_FILES = {
 "0-1":"531-0-1_GLAB_Platform_Check.docx","1-1":"531-1-1_GLAB_Zone_Hardware_ID.docx","1-2":"531-1-2_GLAB_Custody_Walkthrough.docx",
 "1-3":"531-1-3_GLAB_PPE_Hazard_Walkdown.docx","2-1":"531-2-1_GLAB_Asset_Tagging.docx",
 "2-2":"531-2-2_GLAB_RMA_Custody.docx","3-1":"531-3-1_GLAB_DCIM_Records.docx","3-2":"531-3-2_GLAB_Cycle_Count.docx",
 "4-1":"531-4-1_GLAB_Warranty_Claim.docx","4-2":"531-4-2_GLAB_Rack_Stack.docx","5-1":"531-5-1_GLAB_Secure_Decommission.docx","5-2":"531-5-2_GLAB_Disposition_Sanitization.docx"}

SRC_HANDOUT = {n: HANDOUT_FILES[n].replace(f"531-{n}_", f"531-{o}_") for n, o in SRC_ID.items()}   # a renamed lab's handout in the source package

# Old (v2.0) bench objective / kit bullet texts, recognised and dropped when the source is a v2.0 handout
OLD_BULLETS = {
 "Identify real kit items by sight, function and handling requirement at the kit station, scanning each blind card into the tool.",
 "Sequence the lifecycle with physical stage cards and map three stages to the kit item each one reaches for first.",
 "The eight lifecycle stage cards and the kit item cards at your station, DS2208 scanner.",
 "Sort real end-of-life items into disposition bins by scan, witness a destruction and complete the certificate line.",
 "Three physical items at your station (each with a tag or item card and a dossier card), the three bin cards, DS2208 scanner; safety glasses for the crusher demonstration at the instructor bench (SEM Model 0100, instructor-operated).",
}

BENCH_OBJECTIVE = {
 "0-1":"Prove the station works: scan a station card and a scanner card into the lab tool and confirm the ESD kit is present and connected.",
 "1-1":"Identify real kit items by sight, function and handling requirement at the kit station, scanning each blind card into the tool; then sequence the lifecycle with physical stage cards and map three stages to the kit item each one reaches for first.",
 "1-2":"Execute three physical hand-offs of a bagged drive with a complete logbook line at each transfer, mirrored into the tool.",
 "1-3":"Walk a physically staged route, record the genuine hazards without overcalling, and don task-appropriate PPE with a partner fit check.",
 "2-1":"Receive a real carton: weigh it against the ASN, verify contents, print the issued asset tag on the label printer, place it per the standard and scan it back.",
 "2-2":"Pack a failed FRU for return under ESD protection, seal it in a serialized security bag with the serial scanned at sealing, and label the carton with the printed RMA authorization.",
 "3-1":"Find real kit items by scanning their tags against the bench register and record a physical move at the cart, not the desk.",
 "3-2":"Perform a physical blind count of a station bin by scanning every tag, lock it, and reconcile against the register quantity.",
 "4-1":"Read the R640 service tag from the chassis, pack the unit for its path and label it with the printed claim or hold label.",
 "4-2":"Rack a real R640 into the mobile rack with rails, cage nuts and a team lift, dress it, and record the move at the rack.",
 "5-1":"Census real drives, double-bag each one in a serialized security bag with the serial scanned at sealing, and log every bag.",
 "5-2":"Sort real end-of-life items into disposition bins by scan, witness a destruction hand-over and complete the hand-over line.",
}
KIT_BULLETS = {
 "0-1":["Your station kit: DS2208 barcode scanner (USB), ESD mat and wrist strap, the station card and the scanner card taped at the bench."],
 "1-1":["The kit identification station: an R640 server, a hot-swap PSU, a fan module, a drive carrier, a rail kit, cage nuts with a blanking panel, a scanner and a platform cart - each with its model name covered and a blind KIT card beside it.","The eight lifecycle stage cards and the kit item cards at your station, DS2208 barcode scanner."],
 "1-2":["One scrap drive in a sealed anti-static bag carrying a DBD tag, the class chain-of-custody logbook and a pen, DS2208 scanner. Your partner, and the instructor's cage for the third transfer."],
 "1-3":["Your PPE set (hi-vis vest, hard hat, safety glasses, nitrile gloves) and the ESD mat and wrist strap at your station.","The physically staged delivery route in the room (pallet bay, walkway, bench, door) - walked in pairs, nothing touched."],
 "2-1":["The receiving station (shared, in rotation): pallet, bin, bench scale, and your sealed carton with its printed PO/ASN from the mock shipping kit.","Zebra ZD621 label printer at the print station (shared), DS2208 scanner and ESD kit at your station, the FRU item inside your carton."],
 "2-2":["A failed FRU (at a server station: the PSU in the powered-off R640; elsewhere: from the bin), ESD mat and strap, an anti-static shielding bag, one serialized tamper-evident security bag, carton and tape, the custody logbook.","Zebra ZD621 at the print station for the RMA label; DS2208 scanner."],
 "3-1":["The bench register: every R640, FRU, rail kit, cart and the scale in the room carries an asset tag. DS2208 scanner; the platform cart marked STAGE-1; ESD strap for the FRU move."],
 "3-2":["Your station bin (box) of tagged items, DS2208 scanner, an empty tray to set scanned items aside."],
 "4-1":["An R640 at a server station (its pull-out information tag carries the Dell service tag), your FRU, ESD kit, anti-static bag, carton; the ZD621 at the print station for the path label; DS2208 scanner."],
 "4-2":["The PS mobile rack with U-position labels, one R640 with its bench tag, a rail kit, cage nuts, a blanking panel and cable management from the accessories lot, a platform cart, your PPE and ESD strap, DS2208 scanner. Your team-lift partner."],
 "5-1":["The station's chassis bag with its scrap drives (mock DBDs), anti-static shielding bags, serialized tamper-evident security bags, ESD kit, the custody logbook, DS2208 scanner."],
 "5-2":["Three physical items at your station (each with a tag or item card and a dossier card), the three bin cards, DS2208 scanner; the serialized vendor transfer bag for the destruction hand-over at the instructor bench."],
}
DROP_PAT = re.compile(r"(MIR|Barcode scanner|Label printer|Packing bench|Anti-static bag, packing|Tamper-evident bags|Mock DBDs and the manual crusher|PPE kit at your station|The MIR rack, the assigned)", re.I)

def r640(t):
    t = t.replace("R630","R640").replace("thirty-kilogram","twenty-two-kilogram").replace("24 kg","22 kg")
    t = t.replace("MIR (Mobile Infrastructure Rack)", "PS mobile rack (the MIR)")
    return t

def parse(path, strip=False):
    """strip=True: treat the file as a v2.0+ handout and remove its bench additions (bench part, bench objective,
    kit bullets, timing and recording sentences) so the result reads like the v1.2 scenario handout."""
    d = Document(path)
    out = []
    old_kit = {b for v in KIT_BULLETS.values() for b in v} | set(BENCH_OBJECTIVE.values()) | OLD_BULLETS
    for p in d.paragraphs:
        ps = p._element.pPr.find(qn("w:pStyle")) if p._element.pPr is not None else None
        st = {"Heading1":"Heading 1","Heading2":"Heading 2","Heading3":"Heading 3","ListParagraph":"List Paragraph"}.get(ps.get(qn("w:val")) if ps is not None else "", "")
        txt = p.text
        if not txt.strip(): continue
        runs = p.runs
        italic = bool(runs) and all(r.italic for r in runs if r.text.strip())
        if strip:
            if st == "Heading 3" and re.match(r"Part \d+: .*\(at the (bench|rack)\)$", txt): break     # the bench part and everything after it
            if st == "List Paragraph" and txt in old_kit: continue
            txt = re.sub(r" The last \d+ minutes are at the bench \(Part \d+\)\.", "", txt)
            txt = re.sub(r" The bench record \(Part \d+\) is completed during the move, at the rack\.", "", txt)
            txt = txt.replace(" The bench panel inside the tool records the physical part of the lab the same way - scan into it, tick only what is true, and press Build bench record.", "")
            txt = subst(txt)
        if st == "Heading 1": out.append(("h1", txt))
        elif st == "Heading 2": out.append(("h2", txt))
        elif st == "Heading 3": out.append(("h3", txt))
        elif st == "List Paragraph": out.append(("bullet", txt))
        elif txt.startswith("→"): out.append(("arrow", txt[1:].strip()))
        elif italic and "Version" in txt: out.append(("subtitle", txt))
        else: out.append(("body", txt))
    return (subst(d.core_properties.title) if strip else d.core_properties.title), out

def build(lid):
    g, title, slug, sets, les = LAB_META[lid]
    L = LABS[lid]
    src = os.path.join(SRC, "3_Handouts", SRC_HANDOUT.get(lid, HANDOUT_FILES[lid]))
    ctitle, items = parse(src, strip=(os.environ.get("SRC_HAS_BENCH", "1") == "1"))
    files = tool_files_label(lid)
    # count existing questions + steps
    nq = sum(1 for k, t in items if k == "arrow" and ("Answer this as Q" in t or "Attach your image as Q" in t))
    nsteps = max(int(m.group(1)) for k, t in items if k == "body" for m in [re.match(r"Step (\d+)\.", t)] if m)
    # bench steps must continue numbering
    for i, st in enumerate(L["steps"]):
        want = nsteps + 1 + i
        assert st["t"].startswith(f"Step {want}."), (lid, st["t"][:12], want)
    D = Doc(title=ctitle)
    section = None; eq_done = False
    for k, t in items:
        t = r640(t)
        if k == "h1": D.h1(t)
        elif k == "subtitle": D.subtitle(f"Version {VERSION} — Last revised: {VDATE} — {COURSE}")
        elif k == "h2":
            if section == "Equipment / Requirements" and not eq_done:
                for b in KIT_BULLETS[lid]: D.bullet(b); eq_done = True
            section = t; D.h2(t)
        elif k == "h3": D.h3(t)
        elif k == "bullet":
            if section == "Equipment / Requirements" and DROP_PAT.search(t): continue
            if section == "Equipment / Requirements" and "for the physical pass" in t: continue
            if section == "Equipment / Requirements" and "Lab tool:" in t:
                t = t.replace("opened in a browser.", "opened in a browser at your station.")
            D.bullet(t)
        elif k == "arrow": D.arrow(t)
        elif k == "body":
            m = re.search(r"You have (\d+) minutes", t)
            if m:
                add = (f" The last {L['bench_min']} minutes are at the bench ({L['part'].split(':')[0]})." if L["bench_min"]
                       else f" The bench record ({L['part'].split(':')[0]}) is completed during the move, at the rack.")
                t = t.replace(m.group(0), f"You have {L['minutes']} minutes") + add
            if t.startswith("Recording your answers."):
                t += " The bench panel inside the tool records the physical part of the lab the same way - scan into it, tick only what is true, and press Build bench record."
            D.body(t)
    # the bench objective is inserted by the post-pass below
    # ---- append the bench part
    D.h3(L["part"])
    D.body(L["intro"])
    q = nq
    for st in L["steps"]:
        D.body(st["t"])
        if st["cap"] == "rec": D.arrow("Recorded in the lab tool as you work.")
        elif st["cap"] == "text": q += 1; D.arrow(f"Answer this as Q{q} in the lab tool ({files}).")
        elif st["cap"] == "image": q += 1; D.arrow(f"Attach your image as Q{q} in the lab tool ({files}).")
    # insert the bench objective: find the Learning Objectives heading and append after its last bullet
    body = D.d.element.body
    paras = list(D.d.paragraphs)
    idx = None; last_b = None
    for i, p in enumerate(paras):
        ps = p._element.pPr.find(qn("w:pStyle")) if p._element.pPr is not None else None
        sn = ps.get(qn("w:val")) if ps is not None else ""
        if sn == "Heading2" and p.text == "Learning Objectives": idx = i
        elif idx is not None and sn == "ListParagraph": last_b = p
        elif idx is not None and last_b is not None and sn == "Heading2": break
    nb = D.bullet(BENCH_OBJECTIVE[lid]); last_b._element.addnext(nb._element)
    out = os.path.join(OUT, "3_Handouts", HANDOUT_FILES[lid]); D.save(out)
    return nq, q

def main():
    os.makedirs(os.path.join(OUT, "3_Handouts"), exist_ok=True)
    rep = {lid: build(lid) for lid in LAB_ORDER}
    json.dump(rep, open(os.path.join(OUT, "..", "handouts_report.json"), "w"), indent=1)
    print(rep)
if __name__ == "__main__": main()
