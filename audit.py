"""Learner-facing audit: flag internal vocabulary in anything a learner can see."""
import os, re, sys, glob, subprocess, html
from pptx import Presentation
import openpyxl
OUT=os.environ["OUT"]
BAN=[r"\badr\b|adr/",r"\bspec\b",r"Kit Map",r"SET_KEY",r"_FT\b",r"facilitator",r"Instructor_ONLY|instructor-only|instructor only",r"\bseed",r"\bdecoy",r"gen_[a-z]+\.py",r"\bPython\b",r"GitHub",r"generator",r"lockstep",r"curriculum team|curriculum tooling",r"byte-identical",r"harvest",r"BENCH:",r"stations? [0-9], [0-9], [0-9]",r"ATL-009\d\d is (True|False)",r"Set [ABC] prints",r"\bU-17\b|\bU-24\b",r"tool \+ bench",r"answer key|the key\b",r"Not Yet",r"rubric",r"overcall",r"prefix",r"regex|format-checked",r"registry",r"v1\.[0-9]|Version 1\.",r"\bDIAT\b",r"UCI 2225",r"R630"]
def text_of(f):
    if f.endswith((".html",".md")): return open(f,encoding="utf-8").read()
    if f.endswith(".docx"): return subprocess.run(["pandoc",f,"-t","plain","--wrap=none"],capture_output=True,text=True).stdout
    if f.endswith(".pptx"):
        p=Presentation(f); return "\n".join(sh.text_frame.text for s in p.slides for sh in s.shapes if sh.has_text_frame)  # slide text only (notes are instructor-facing)
    if f.endswith(".xlsx"):
        wb=openpyxl.load_workbook(f); return "\n".join(str(c.value) for row in wb["Questions"].iter_rows() for c in row if c.value)
    return ""
files=sorted(glob.glob(OUT+"/3_Handouts/*")+glob.glob(OUT+"/4_Lab_Tools/*")+glob.glob(OUT+"/2_Decks/*.pptx")+glob.glob(OUT+"/6_Assessments/*.xlsx"))
hits=0
for f in files:
    t=text_of(f)
    if f.endswith(".html"):  # visible text only: drop scripts/styles/comments, unescape
        t=re.sub(r"<script[\s\S]*?</script>|<style[\s\S]*?</style>|<!--[\s\S]*?-->","",t); t=html.unescape(re.sub(r"<[^>]+>"," ",t))
    for pat in BAN:
        for m in re.finditer(pat,t,re.I):
            hits+=1; ctx=t[max(0,m.start()-50):m.end()+50].replace("\n"," ")
            print(f"{os.path.relpath(f,OUT)} | {pat} | …{ctx}…")
print("HITS",hits)
